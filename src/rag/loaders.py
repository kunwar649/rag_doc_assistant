"""Document loaders: PDF, DOCX and TXT/MD -> list[Page].

Each loader keeps source metadata (file name, page) so answers can cite where
they came from later.
"""
import logging
import re
from collections import Counter
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from .models import Page

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def _clean(text: str) -> str:
    # Collapse runs of whitespace that PDF extraction often produces.
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def _line_key(line: str) -> str:
    # "Page 7" and "Page 8" should count as the same running footer.
    return re.sub(r"\d+", "#", line.lower()).strip()


def strip_running_headers(pages: list[Page], edge_lines: int = 2, min_ratio: float = 0.5) -> list[Page]:
    """Remove headers/footers repeated at the top or bottom of most pages.

    Lines like "Tennessee License Validation Adapter | Technical Walkthrough"
    and "Page 7" otherwise end up in every chunk: they add no information and
    make unrelated chunks look similar to each other.
    Only the first/last `edge_lines` of each page are considered, so the same
    sentence appearing in the body is never removed.
    """
    if len(pages) < 3:  # too few pages to tell a header from content
        return pages

    def edge(n_lines: int) -> int:
        # On short pages the top and bottom windows would overlap the body; shrink them.
        return min(edge_lines, n_lines // 3)

    counts: Counter[str] = Counter()
    for p in pages:
        lines = p.text.splitlines()
        e = edge(len(lines))
        if e:
            counts.update({_line_key(line) for line in lines[:e] + lines[-e:]})
    repeated = {k for k, c in counts.items() if k and c / len(pages) >= min_ratio}
    if not repeated:
        return pages

    cleaned = []
    for p in pages:
        lines = p.text.splitlines()
        n, e = len(lines), edge(len(lines))
        kept = [
            line
            for i, line in enumerate(lines)
            if not ((i < e or i >= n - e) and _line_key(line) in repeated)
        ]
        # Never drop a page entirely because of this heuristic.
        cleaned.append(Page(text="\n".join(kept), source=p.source, page=p.page) if kept else p)
    return cleaned


def load_pdf(path: Path) -> list[Page]:
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        if text:  # skip blank / image-only pages
            pages.append(Page(text=text, source=path.name, page=i))
    if not pages:
        logger.warning("No extractable text in %s (scanned PDF? needs OCR)", path.name)
    return strip_running_headers(pages)


def load_docx(path: Path) -> list[Page]:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Tables hold a lot of real-world content; flatten them row by row.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = _clean("\n".join(parts))
    return [Page(text=text, source=path.name)] if text else []


def load_text(path: Path) -> list[Page]:
    text = _clean(path.read_text(encoding="utf-8", errors="ignore"))
    return [Page(text=text, source=path.name)] if text else []


_LOADERS = {".pdf": load_pdf, ".docx": load_docx, ".txt": load_text, ".md": load_text}


def load_file(path: Path) -> list[Page]:
    ext = path.suffix.lower()
    if ext not in _LOADERS:
        raise ValueError(f"Unsupported file type: {ext}")
    return _LOADERS[ext](path)


def load_directory(directory: Path) -> list[Page]:
    pages: list[Page] = []
    for path in sorted(directory.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            try:
                file_pages = load_file(path)
                logger.info("Loaded %s: %d page(s)", path.name, len(file_pages))
                pages.extend(file_pages)
            except Exception:  # one bad file shouldn't stop the whole ingest
                logger.exception("Failed to load %s", path.name)
    return pages
