import pytest
from docx import Document

from rag.loaders import load_directory, load_file


def test_load_txt(tmp_path):
    f = tmp_path / "notes.txt"
    f.write_text("Line   one\n\n\nLine two   ")
    pages = load_file(f)
    assert len(pages) == 1
    assert pages[0].text == "Line one\nLine two"
    assert pages[0].page is None


def test_load_docx_includes_tables(tmp_path):
    f = tmp_path / "spec.docx"
    doc = Document()
    doc.add_paragraph("Intro paragraph")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Plan"
    table.rows[0].cells[1].text = "Free"
    doc.save(f)
    text = load_file(f)[0].text
    assert "Intro paragraph" in text
    assert "Plan | Free" in text


def test_unsupported_extension(tmp_path):
    f = tmp_path / "image.png"
    f.write_bytes(b"x")
    with pytest.raises(ValueError):
        load_file(f)


def test_strip_running_headers():
    from rag.loaders import strip_running_headers
    from rag.models import Page

    topics = ["Setup steps", "Error handling", "CSV output", "CAPTCHA flow"]
    pages = [Page("Cover page\nIntro text", "d.pdf", 1)] + [
        Page(
            f"My Report | Walkthrough\nPage {i}\n{topic}\nDetails about {topic.lower()}.\n"
            f"My Report | Walkthrough\nMore on {topic.lower()}.\nClosing note on {topic.lower()}.",
            "d.pdf",
            i,
        )
        for i, topic in enumerate(topics, start=2)
    ]
    cleaned = strip_running_headers(pages)
    assert cleaned[0].text == "Cover page\nIntro text"
    lines = cleaned[1].text.splitlines()
    assert lines[0] == "Setup steps"  # header + page number removed
    assert "My Report | Walkthrough" in lines  # same text in the body is kept
    assert lines[-1] == "Closing note on setup steps."
    assert all("Page " not in p.text for p in cleaned[1:])


def test_strip_running_headers_never_empties_a_page():
    from rag.loaders import strip_running_headers
    from rag.models import Page

    # Pages made only of the same repeated lines, at several lengths: none may end up empty.
    for n_lines in range(1, 8):
        pages = [Page("\n".join(["Header", "Page 1"] * n_lines), "d.pdf", i) for i in range(1, 5)]
        assert all(p.text.strip() for p in strip_running_headers(pages))


def test_strip_running_headers_leaves_short_docs_alone():
    from rag.loaders import strip_running_headers
    from rag.models import Page

    pages = [Page("Header\nA", "d.pdf", 1), Page("Header\nB", "d.pdf", 2)]
    assert strip_running_headers(pages) == pages


def test_directory_skips_unsupported_and_empty(tmp_path):
    (tmp_path / "a.txt").write_text("content")
    (tmp_path / "empty.txt").write_text("   ")
    (tmp_path / "skip.csv").write_text("a,b")
    pages = load_directory(tmp_path)
    assert [p.source for p in pages] == ["a.txt"]
